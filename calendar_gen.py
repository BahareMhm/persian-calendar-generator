import requests
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Border, Side, Alignment, Font
from openpyxl.worksheet.page import PageMargins


API_URL = "https://pnldev.com/api/calender"

NOTE_ROWS = 4
YEAR = 1406

MONTH_NAMES = [
    "فروردین", "اردیبهشت", "خرداد",
    "تیر","مرداد","شهریور",
    "مهر", "آبان", "آذر",
    "دی","بهمن","اسفند",
]

WEEKDAYS = ["شنبه","یکشنبه","دوشنبه","سه‌شنبه","چهارشنبه","پنجشنبه","جمعه"]

def get_calendar_from_api(year):

    response = requests.get(API_URL, params={"year": year}, timeout=20)

    response.raise_for_status()

    data = response.json()

    if data.get("status") is not True:
        raise ValueError("API status is not true.")

    return data["result"]


def parse_api_data(result):
    """
    change api structur
    output:
    {
    (ماه, روز): {"weekday": ..., "holiday": ..., "events": [...]}
    }
    """

    days = {}

    for month_key, month_data in result.items():

        for day_key, day_data in month_data.items():

            solar = day_data["solar"]
            month = solar["month"]
            day = solar["day"]

            days[(month, day)] = {
                "weekday": solar.get("dayWeek"),
                "holiday": day_data.get("holiday", False),
                "events": day_data.get("event", [])
            }

    return days


def weekday_to_index(day_week):
    """
    output:
        0 = شنبه
        1 = یکشنبه
        2 = دوشنبه
        ...
        6 = جمعه
    """

    mapping = {"ش": 0, "ی": 1, "د": 2, "س": 3, "چ": 4, "پ": 5, "ج": 6}

    return mapping.get(day_week)


def create_workbook():

    wb = Workbook()
    ws = wb.active

    ws.title = f"Calendar {YEAR}"
    ws.sheet_view.rightToLeft = True

    return wb, ws


def create_styles():

    orange_fill = PatternFill(fill_type="solid", fgColor="F4B400")
    turquoise_fill = PatternFill(fill_type="solid", fgColor="46BDC6")
    light_yellow = PatternFill(fill_type="solid", fgColor="FFF2CC")
    light_blue = PatternFill(fill_type="solid", fgColor="D9EAF7")
    header_fill = PatternFill(fill_type="solid", fgColor="9DB963")

    thin_side = Side(style="thin", color="FFFFFF")
    border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    normal_font = Font(name="Tahoma", size=10)
    bold_font = Font(name="Tahoma", size=11, bold=True)
    holiday_font = Font(name="Tahoma", size=10, bold=True, color="9C0006")

    return {
        "orange_fill": orange_fill,
        "header_fill": header_fill,
        "turquoise_fill": turquoise_fill,
        "light_yellow": light_yellow,
        "light_blue": light_blue,
        "border": border,
        "normal_font": normal_font,
        "bold_font": bold_font,
        "holiday_font": holiday_font,
    }


def get_month_fill(month, styles):

    if month % 2 == 1:
        return styles["turquoise_fill"]

    return styles["orange_fill"]


def create_header(ws, styles):

    #column A = month
    ws.cell(row=1, column=1).value = "نام ماه"
    ws.cell(row=1, column=1).border = styles["border"]
    ws.cell(row=1, column=1).font = styles["bold_font"]
    ws.cell(row=1, column=1).alignment = Alignment(horizontal="center", vertical="center")

    #column B to H = weekdays
    for i, weekday in enumerate(WEEKDAYS, start=2):

        cell = ws.cell(row=1, column=i)
        cell.value = weekday
        cell.fill = styles["header_fill"]
        cell.border = styles["border"]
        cell.font = styles["bold_font"]
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    ws.row_dimensions[1].height = 25


def get_note_fill(weekday_index, styles):

    if weekday_index % 2 == 0:
        return styles["light_yellow"]

    return styles["light_blue"]


def write_week(ws, start_row, week_days, days_data, styles):

    date_row = start_row

    months_in_week = []

    for item in week_days:

        if item is not None:
            month, day = item

            if month not in months_in_week:
                months_in_week.append(month)

 
    if months_in_week:

        second_month = months_in_week[-1]

        month_cell = ws.cell(row=date_row, column=1)
        month_cell.value = MONTH_NAMES[second_month - 1]

        month_cell.fill = get_month_fill(month, styles)
        month_cell.border = styles["border"]
        month_cell.font = styles["bold_font"]
        month_cell.alignment = Alignment(horizontal="center", vertical="center")

    for weekday_index, item in enumerate(week_days):

        column = weekday_index + 2

        date_cell = ws.cell(row=date_row, column=column)

        date_cell.border = styles["border"]
        date_cell.alignment = Alignment(horizontal="center", vertical="center")
        date_cell.font = styles["bold_font"]

        if item is not None:

            month, day = item
            date_cell.value = day
            date_cell.fill = get_month_fill(month, styles)
            day_info = days_data[(month, day)]

            if day_info["holiday"]:
                date_cell.font = styles["holiday_font"]

        else:
            date_cell.fill = styles["turquoise_fill"]



        note_fill = get_note_fill(weekday_index, styles)

        for i in range(1, NOTE_ROWS + 1):

            note_row = date_row + i
            note_cell = ws.cell(row=note_row, column=column)
            note_cell.fill = note_fill
            note_cell.border = styles["border"]
            note_cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
            note_cell.font = styles["normal_font"]

            #calendar events
            if (item is not None and i == 1):

                month, day = item
                events = days_data[(month, day)]["events"]

                if events:

                    note_cell.value = "\n".join(events)
                    if days_data[(month, day)]["holiday"]:
                        note_cell.font = styles["holiday_font"]


    #cells height
    ws.row_dimensions[date_row].height = 24
    for i in range(1, NOTE_ROWS + 1):
        ws.row_dimensions[date_row + i].height = 30



def create_calendar(ws, days_data, styles):
    """
    create all weeks in ws
    """
    current_row = 2
    current_week = [None] * 7

    all_days = sorted(days_data.keys(), key=lambda x: (x[0], x[1]))

    for month, day in all_days:

        info = days_data[(month, day)]
        weekday_index = weekday_to_index(info["weekday"])

        if weekday_index is None:
            raise ValueError(f"not a valid weekday: {info['weekday']}")

   
        current_week[weekday_index] = (month, day)

        # اگر جمعه است، هفته کامل شده
        if weekday_index == 6:

            write_week(ws, current_row, current_week, days_data, styles)
            current_row += NOTE_ROWS + 1
            current_week = [None] * 7

    # هفته ناقص آخر سال
    if any(item is not None for item in current_week):
        write_week(ws, current_row, current_week, days_data, styles)


def configure_sheet(ws):

    #set month column width
    ws.column_dimensions["A"].width = 12

    #set days column width
    for column in range(2, 9):

        letter = chr(64 + column)
        ws.column_dimensions[letter].width = 18

    #view rightToLeft
    ws.sheet_view.rightToLeft = True

    #freeze header in scroll
    ws.freeze_panes = "B2"

    #delete Gridline
    ws.sheet_view.showGridLines = False

    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = (ws.PAPERSIZE_A4)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.4, bottom=0.4, header=0.2, footer=0.2)


    ws.print_title_rows = "1:1"
    ws.print_area = (f"A1:H{ws.max_row}")


#==================================================================================


def main():

    result = get_calendar_from_api(YEAR)
    print("got the data from API.")

    days_data = parse_api_data(result)
    print(f"Number of days that been resived: {len(days_data)}")

    #Excel
    wb, ws = create_workbook()
    styles = create_styles()

    create_header(ws, styles)

    create_calendar(ws, days_data, styles)

    configure_sheet(ws)

  
    filename = (f"calendar_{YEAR}.xlsx")
    wb.save(filename)
    print(f"\ncalendar {YEAR} successfully savesd. {(filename)}")


if __name__ == "__main__":
    main()
    